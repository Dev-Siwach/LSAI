import os
from pathlib import Path
from typing import Dict, Any, List, Optional, Union
from datetime import datetime

from config.settings import get_settings


class DeliverableGenerator:
    """Deliverable generation tool for .docx, .xlsx, and .pptx.

    Uses lazy imports so the tool can be safely loaded and tested in air-gapped
    or lightweight environments before external document libraries are installed.
    """

    def __init__(self):
        self.settings = get_settings()
        self.deliverables_dir = self.settings.DELIVERABLES_DIR
        self.deliverables_dir.mkdir(parents=True, exist_ok=True)

    def _generate_filepath(self, prefix: str, ext: str) -> Path:
        """Generate a unique filepath for a new deliverable."""
        self.deliverables_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{prefix}_{timestamp}{ext}"
        return self.deliverables_dir / filename

    def generate_docx_approval_note(self, data: Dict[str, Any]) -> str:
        """Generate a formatted PSU/government .docx note sheet.

        data should contain:
            ref_number (str): Reference or file tracking ID
            subject (str): Official subject line
            findings (str): Technical or operational observations
            recommendations (str): Proposed action and hierarchy sign-off
            annexures (str, optional): Attached documents list
            tabular_data (list of dicts or list of lists, optional): Table data
        """
        try:
            import docx
            from docx.shared import Pt, Inches, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.enum.table import WD_TABLE_ALIGNMENT
        except ImportError:
            raise ImportError("python-docx is required for docx generation. Please install python-docx.")

        doc = docx.Document()

        # Set standard margins (1 inch)
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # Header
        header = doc.add_heading("SOVEREIGN INDUSTRIAL WORKBENCH — APPROVAL NOTE", level=0)
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Ref & Date paragraph
        ref_para = doc.add_paragraph()
        ref_run = ref_para.add_run("Ref No: ")
        ref_run.bold = True
        ref_para.add_run(str(data.get("ref_number", "PSU/REF/2026/001")) + "\t\t\t\t")
        date_run = ref_para.add_run("Date: ")
        date_run.bold = True
        ref_para.add_run(datetime.now().strftime("%d-%b-%Y"))

        divider = doc.add_paragraph("—" * 60)
        divider.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Subject
        subj_para = doc.add_paragraph()
        s_label = subj_para.add_run("Subject: ")
        s_label.bold = True
        subj_para.add_run(str(data.get("subject", "Industrial Evaluation and Operational Recommendation")))

        # 1. Findings
        doc.add_heading("1. Technical Findings & Analysis", level=1)
        doc.add_paragraph(str(data.get("findings", "No technical findings provided.")))

        # Tabular Data Annexure
        table_data = data.get("tabular_data")
        if table_data and isinstance(table_data, list) and len(table_data) > 0:
            doc.add_heading("Tabular Data Annexure", level=2)
            if isinstance(table_data[0], dict):
                headers = list(table_data[0].keys())
                table = doc.add_table(rows=1, cols=len(headers))
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for i, col_name in enumerate(headers):
                    hdr_cell = table.rows[0].cells[i]
                    hdr_cell.text = str(col_name)
                    for paragraph in hdr_cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
                for row_dict in table_data:
                    row_cells = table.add_row().cells
                    for i, col_name in enumerate(headers):
                        row_cells[i].text = str(row_dict.get(col_name, ""))
            elif isinstance(table_data[0], (list, tuple)):
                headers = [str(c) for c in table_data[0]]
                table = doc.add_table(rows=1, cols=len(headers))
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for i, col_name in enumerate(headers):
                    hdr_cell = table.rows[0].cells[i]
                    hdr_cell.text = str(col_name)
                    for paragraph in hdr_cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
                for row_list in table_data[1:]:
                    row_cells = table.add_row().cells
                    for i, val in enumerate(row_list):
                        if i < len(row_cells):
                            row_cells[i].text = str(val)

        # 2. Recommendations
        doc.add_heading("2. Recommendations & Approval Hierarchy", level=1)
        doc.add_paragraph(str(data.get("recommendations", "No recommendations provided.")))

        # 3. Annexures / References
        doc.add_heading("3. References & Annexures", level=1)
        doc.add_paragraph(str(data.get("annexures", "None.")))

        # Sign-off block
        doc.add_paragraph("\n\nSubmitted for approval,\n\n_______________________\nLead Competent Authority")

        filepath = self._generate_filepath("Approval_Note", ".docx")
        doc.save(str(filepath))
        return str(filepath)

    def generate_excel_calculation_sheet(self, data: Dict[str, Any]) -> str:
        """Generate .xlsx workbook with calculation rows and styled headers.

        data should contain:
            title (str, optional): Sheet title
            headers (list of str): Column header strings
            rows (list of lists): Row values
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            raise ImportError("openpyxl is required for Excel generation. Please install openpyxl.")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Calculations"

        title = str(data.get("title", "Engineering Calculation Sheet"))
        ws.append([title])
        title_cell = ws["A1"]
        title_cell.font = Font(size=14, bold=True, color="1F497D")

        ws.append([])  # Empty row spacer

        headers = data.get("headers", [])
        if headers:
            ws.append(headers)
            header_row_idx = 3
            fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            font = Font(bold=True, color="FFFFFF")
            align = Alignment(horizontal="center", vertical="center")

            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=header_row_idx, column=col)
                cell.font = font
                cell.fill = fill
                cell.alignment = align

        thin_border = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3"),
        )

        for row_data in data.get("rows", []):
            ws.append(list(row_data))
            for cell in ws[ws.max_row]:
                cell.border = thin_border

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        filepath = self._generate_filepath("Calc_Sheet", ".xlsx")
        wb.save(str(filepath))
        wb.close()
        return str(filepath)

    def generate_pptx_briefing(self, data: Dict[str, Any]) -> str:
        """Generate .pptx summary slides for executive briefing.

        data should contain:
            title (str): Presentation title
            subtitle (str, optional): Subtitle or date
            slides (list of dicts): Each dict with 'title' (str) and 'content' (str or list)
        """
        try:
            import pptx
            from pptx.util import Inches as PptxInches, Pt as PptxPt
        except ImportError:
            raise ImportError("python-pptx is required for PowerPoint generation. Please install python-pptx.")

        prs = pptx.Presentation()

        # Title Slide
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]

        title.text = str(data.get("title", "Executive Technical Briefing"))
        subtitle.text = str(data.get("subtitle", f"Generated on {datetime.now().strftime('%Y-%m-%d')}"))

        # Bullet / Content Slides
        bullet_slide_layout = prs.slide_layouts[1]
        for slide_data in data.get("slides", []):
            slide = prs.slides.add_slide(bullet_slide_layout)
            shapes = slide.shapes
            title_shape = shapes.title
            body_shape = shapes.placeholders[1]

            title_shape.text = str(slide_data.get("title", "Summary"))

            tf = body_shape.text_frame
            content = slide_data.get("content", [])
            if isinstance(content, list):
                for i, item in enumerate(content):
                    if i == 0:
                        tf.text = str(item)
                    else:
                        p = tf.add_paragraph()
                        p.text = str(item)
                        p.level = 0
            else:
                tf.text = str(content)

        filepath = self._generate_filepath("Briefing", ".pptx")
        prs.save(str(filepath))
        return str(filepath)
